"""DOCX parser implementation using python-docx.

Purpose:
    Extract text from paragraphs and tables in DOCX documents in a stateless, thread-safe manner.
"""

from __future__ import annotations

import os
from pathlib import Path
import docx

from ats_engine.domain.document_processing.exceptions import CorruptedDocumentError, DocumentReadError
from ats_engine.domain.document_processing.models import RawDocument
from ats_engine.domain.document_processing.parser import DocumentParser


class DocxDocumentParser(DocumentParser):
    """Stateless python-docx-based DOCX text extractor."""

    def parse(self, file_path: Path) -> RawDocument:
        """Parse raw content from a DOCX file.

        Args:
            file_path: Absolute path to the DOCX file.
        """
        if not file_path.exists():
            raise DocumentReadError(f"File not found: {file_path}")

        try:
            file_size = os.path.getsize(file_path)
        except OSError as error:
            raise DocumentReadError(f"Failed to read file size for {file_path}: {error}") from error

        try:
            doc = docx.Document(file_path)
        except Exception as error:
            raise CorruptedDocumentError(f"Failed to open DOCX document {file_path.name}: {error}") from error

        try:
            segments: list[str] = []
            
            # 1. Extract standard paragraphs
            for paragraph in doc.paragraphs:
                segments.append(paragraph.text)
            
            # 2. Extract cell texts from all tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text = cell.text.strip()
                        if text:
                            segments.append(text)
            
            raw_content = "\n".join(segments)
            
            # Since DOCX layout rendering is application-specific, page limits
            # are not stored in standard XML. We treat the document as a single page range.
            return RawDocument(
                filename=file_path.name,
                file_size_bytes=file_size,
                raw_content=raw_content,
                pages=(raw_content,),
            )
        except Exception as error:
            raise CorruptedDocumentError(f"Failed to extract text from DOCX {file_path.name}: {error}") from error
