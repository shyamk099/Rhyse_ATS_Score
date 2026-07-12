"""Unit tests for the Document Processing Foundation.

Purpose:
    Verify PDF and DOCX parsing, text normalization, signature-based factory routing,
    metadata audits, and error boundary handling.
"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
import docx
import fitz

from ats_engine.domain.document_processing.exceptions import (
    CorruptedDocumentError,
    DocumentReadError,
    UnsupportedFormatError,
)
from ats_engine.domain.document_processing.docx_parser import DocxDocumentParser
from ats_engine.domain.document_processing.factory import DocumentParserFactory
from ats_engine.domain.document_processing.models import ParsingResult
from ats_engine.domain.document_processing.normalization import TextNormalizer
from ats_engine.domain.document_processing.pdf_parser import PdfDocumentParser
from ats_engine.domain.document_processing.registry import DocumentParserRegistry
from ats_engine.domain.document_processing.service import DocumentProcessingService


class DocumentProcessingTests(unittest.TestCase):
    """Test suite validating stateless parsing and the cleaning pipeline."""

    def setUp(self) -> None:
        """Set up temporary directories and registry dependencies."""
        self._temp_dir = tempfile.TemporaryDirectory()
        self._base_dir = Path(self._temp_dir.name)

        # Wire the parser registry and factory
        registry = DocumentParserRegistry()
        registry.register("pdf", PdfDocumentParser)
        registry.register("docx", DocxDocumentParser)

        self._factory = DocumentParserFactory(registry)
        self._service = DocumentProcessingService(self._factory)

    def tearDown(self) -> None:
        """Release temporary directories after test run."""
        self._temp_dir.cleanup()

    def test_parses_valid_pdf_file(self) -> None:
        """Service resolves PDF format and parses its content successfully."""
        pdf_path = self._base_dir / "test.pdf"
        self._create_pdf(pdf_path, ["Page 1 text.", "Page 2 text here."])

        result = self._service.parse(pdf_path)

        self.assertIsInstance(result, ParsingResult)
        self.assertEqual("test.pdf", result.raw_document.filename)
        self.assertEqual(2, result.metadata.page_count)
        self.assertIn("Page 1 text.", result.normalized_document.cleaned_content)
        self.assertIn("Page 2 text here.", result.normalized_document.cleaned_content)
        self.assertEqual("PdfDocumentParser", result.metadata.parser_used)

    def test_parses_valid_docx_file_including_tables(self) -> None:
        """Service resolves DOCX format and extracts paragraph and table contents."""
        docx_path = self._base_dir / "test.docx"
        self._create_docx(
            docx_path,
            paragraphs=["Paragraph 1 content.", "Paragraph 2 content."],
            table_cells=["Table cell text 1", "Table cell text 2"],
        )

        result = self._service.parse(docx_path)

        self.assertEqual("test.docx", result.raw_document.filename)
        self.assertEqual(1, result.metadata.page_count)  # Treat docx as 1 page range
        self.assertIn("Paragraph 1 content.", result.normalized_document.cleaned_content)
        self.assertIn("Table cell text 1", result.normalized_document.cleaned_content)
        self.assertEqual("DocxDocumentParser", result.metadata.parser_used)

    def test_normalizer_pipeline_cleans_text(self) -> None:
        """TextNormalizer collapses spaces, standardizes line breaks, and repairs hyphens."""
        dirty_text = (
            "Unicode test: \u2014 em dash.  Extra   spaces  \n"
            "word-\n    wrap hyphenation.\n\n\nEmpty lines above."
        )

        clean_text = TextNormalizer.normalize(dirty_text)

        # Em dash NFKC normalizes to standard unicode, spaces collapse
        self.assertIn("Unicode test: — em dash. Extra spaces", clean_text)
        # Word hyphen split at newline is repaired
        self.assertIn("wordwrap hyphenation.", clean_text)
        # Consecutive empty lines collapse
        self.assertNotIn("\n\n\n", clean_text)

    def test_fails_on_missing_file(self) -> None:
        """Parsing a non-existent path raises DocumentReadError."""
        with self.assertRaises(DocumentReadError):
            self._service.parse(self._base_dir / "nonexistent.pdf")

    def test_fails_on_unsupported_magic_bytes(self) -> None:
        """Files with unsupported header magic bytes raise UnsupportedFormatError."""
        txt_path = self._base_dir / "test.txt"
        txt_path.write_text("Plain text content. Not a PDF or DOCX.", encoding="utf-8")

        with self.assertRaises(UnsupportedFormatError):
            self._service.parse(txt_path)

    def test_fails_on_corrupt_pdf(self) -> None:
        """Corrupted PDF files raise CorruptedDocumentError."""
        corrupt_path = self._base_dir / "corrupt.pdf"
        # Write PDF signature but followed by garbage binary content
        corrupt_path.write_bytes(b"%PDF-1.4\nGarbage data\x00\xff")

        with self.assertRaises(CorruptedDocumentError):
            self._service.parse(corrupt_path)

    def test_fails_on_corrupt_docx(self) -> None:
        """Corrupted DOCX files raise CorruptedDocumentError."""
        corrupt_path = self._base_dir / "corrupt.docx"
        # Write ZIP signature but invalid archive payload
        corrupt_path.write_bytes(b"PK\x03\x04\x00\x00\x00\x00")

        with self.assertRaises(CorruptedDocumentError):
            self._service.parse(corrupt_path)

    def test_parses_empty_pdf(self) -> None:
        """Empty text-layer PDF parses successfully returning empty clean text."""
        pdf_path = self._base_dir / "empty.pdf"
        self._create_pdf(pdf_path, ["", ""])

        result = self._service.parse(pdf_path)
        self.assertEqual("", result.normalized_document.cleaned_content)
        self.assertEqual(0, result.metadata.character_count)
        self.assertEqual(0, result.metadata.paragraph_count)

    def _create_pdf(self, path: Path, pages: list[str]) -> None:
        doc = fitz.open()
        for text in pages:
            page = doc.new_page()
            if text:
                page.insert_text((50, 50), text)
        doc.save(path)
        doc.close()

    def _create_docx(
        self,
        path: Path,
        paragraphs: list[str],
        table_cells: list[str] | None = None,
    ) -> None:
        doc = docx.Document()
        for p_text in paragraphs:
            doc.add_paragraph(p_text)
        
        if table_cells:
            table = doc.add_table(rows=1, cols=len(table_cells))
            row = table.rows[0]
            for idx, cell_text in enumerate(table_cells):
                row.cells[idx].text = cell_text
                
        doc.save(path)
