"""Utility to generate sample files for integration testing and pipeline validation.

Purpose:
    Build PDF, DOCX, empty, and corrupted files under samples/.
"""

from __future__ import annotations

import os
from pathlib import Path
import docx
import fitz


def main() -> None:
    """Generate sample files under samples/."""
    samples_dir = Path("samples")
    samples_dir.mkdir(parents=True, exist_ok=True)

    # 1. Valid PDF Resume
    pdf_text_pages = [
        "John Doe\nEmail: john.doe@example.com\nPhone: +1-555-123-4567\nLinkedIn: linkedin.com/in/johndoe\nGitHub: github.com/johndoe",
        "SUMMARY\nExperienced software engineer with a track record of building python applications.",
        "SKILLS\nPython, Docker, AWS Cloud, Git",
        "EXPERIENCE\nSoftware Engineer\nAcme Corp\n2020 - 2022\nBuilt python web scrapers and crawlers using Docker.",
        "EDUCATION\nBS in Computer Science\nMIT\n2016 - 2020",
        "PROJECTS\nWeb Crawler\nRepo: https://github.com/johndoe/crawler\nDemo: https://crawler-demo.com\nImplemented in Python and Docker.",
        "CERTIFICATIONS\nAWS Certified Solutions Architect\nTech: in Docker\nCredential URL: https://aws.amazon.com/verify/123",
    ]
    doc = fitz.open()
    for text in pdf_text_pages:
        page = doc.new_page()
        page.insert_text((50, 50), text)
    doc.save(samples_dir / "resume.pdf")
    doc.close()
    print("Generated samples/resume.pdf")

    # 2. Valid DOCX Resume
    doc_docx = docx.Document()
    for text in pdf_text_pages:
        doc_docx.add_paragraph(text)
    doc_docx.save(samples_dir / "resume.docx")
    print("Generated samples/resume.docx")

    # 3. Empty Resume (No content)
    doc_empty = fitz.open()
    doc_empty.new_page()
    doc_empty.save(samples_dir / "empty.pdf")
    doc_empty.close()
    print("Generated samples/empty.pdf")

    doc_docx_empty = docx.Document()
    doc_docx_empty.save(samples_dir / "empty.docx")
    print("Generated samples/empty.docx")

    # 4. Corrupted PDF
    with open(samples_dir / "corrupted.pdf", "wb") as f:
        f.write(b"this is not a valid pdf header %PDF-1.4 corrupt content here")
    print("Generated samples/corrupted.pdf")

    # 5. Unsupported file format
    with open(samples_dir / "unsupported.txt", "w") as f:
        f.write("unsupported text file content")
    print("Generated samples/unsupported.txt")


if __name__ == "__main__":
    main()
